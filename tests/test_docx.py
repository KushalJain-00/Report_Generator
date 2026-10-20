"""Converter check: run with `python tests/test_docx.py`."""
import sys
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document

from app import md_to_docx

SAMPLE = """# Project Title

Intro paragraph with **bold text**, *italic*, and `inline code`.

## Scope

- First bullet
- Second bullet with **emphasis**

1. Ordered one
2. Ordered two

> A blockquote line

| Phase | Duration | Owner |
|-------|----------|-------|
| Discovery | 2 weeks | Alice |
| Delivery | 4 weeks | Bob |

```python
def hello():
    return "code block"
```

### Closing

Final paragraph.
"""


def main():
    data = md_to_docx(SAMPLE)
    assert data[:2] == b"PK", "docx is a zip container"

    doc = Document(BytesIO(data))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    texts = [p.text for p in paras]

    assert any(p.style.name == "Heading 1" and p.text == "Project Title" for p in paras), "H1"
    assert any(p.style.name == "Heading 2" and p.text == "Scope" for p in paras), "H2"
    assert any(p.style.name == "Heading 3" and p.text == "Closing" for p in paras), "H3"

    bold_runs = [r for p in paras for r in p.runs if r.bold]
    assert any("bold text" in r.text for r in bold_runs), "bold run"
    italic_runs = [r for p in paras for r in p.runs if r.italic]
    assert any("italic" in r.text for r in italic_runs), "italic run"

    bullets = [p for p in paras if p.style.name == "List Bullet"]
    assert len(bullets) == 2, f"bullets: {len(bullets)}"
    numbers = [p for p in paras if p.style.name == "List Number"]
    assert len(numbers) == 2, f"numbers: {len(numbers)}"

    quotes = [p for p in paras if p.style.name == "Quote"]
    assert quotes and quotes[0].text == "A blockquote line", "quote"

    assert len(doc.tables) == 1, "one table"
    t = doc.tables[0]
    assert len(t.rows) == 3 and len(t.columns) == 3, f"table shape {len(t.rows)}x{len(t.columns)}"
    assert t.rows[0].cells[0].text == "Phase"
    assert t.rows[1].cells[1].text == "2 weeks"
    assert all(r.bold for r in t.rows[0].cells[0].paragraphs[0].runs), "header bold"

    assert any("def hello():" in t2 or "def hello():" in t2 for t2 in texts) or any(
        "def hello():" in r.text for p in paras for r in p.runs
    ), "code block content preserved"

    assert any("Final paragraph." in t2 for t2 in texts), "closing paragraph"

    # edge: unterminated code fence still yields content
    doc2 = Document(BytesIO(md_to_docx("# T\n```\ncode here")))
    assert any("code here" in r.text for p in doc2.paragraphs for r in p.runs), "unterminated fence"

    # edge: empty input -> valid docx
    Document(BytesIO(md_to_docx("")))

    print("test_docx: all asserts passed")


if __name__ == "__main__":
    main()
