"""Rendering — markdown to styled HTML (preview/PDF), PDF, DOCX.

Pure sync functions; callers wrap with asyncio.to_thread. Ported from
legacy app.py (markdown_to_html / generate_preview_html merged into one
build_html with a pdf flag).
"""

import time
from io import BytesIO

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from weasyprint import HTML as WeasyHTML

_PDF_STYLE = """
@page {
    size: A4;
    margin: 2.5cm 2cm;
    @top-center { content: "__TITLE__"; font-size: 9px; color: #888; }
    @bottom-center { content: "Page " counter(page) " of " counter(pages); font-size: 9px; color: #888; }
}
body{font-family:'Segoe UI','Helvetica Neue',Arial,sans-serif;color:#1a1a2e;line-height:1.7;font-size:13px;max-width:100%}
h1{font-size:22px;border-bottom:3px solid #6366f1;padding-bottom:8px;margin-top:0;color:#1a1a2e}
h2{font-size:17px;color:#2d2d44;margin-top:28px;border-left:4px solid #6366f1;padding-left:10px}
h3{font-size:14px;color:#3a3a52;margin-top:20px}
table{width:100%;border-collapse:collapse;margin:14px 0;font-size:12px}
th{background:#f0f0f5;padding:8px 10px;text-align:left;border:1px solid #d0d0e0;font-weight:600}
td{padding:6px 10px;border:1px solid #d0d0e0}
tr:nth-child(even){background:#f8f8fc}
ul,ol{padding-left:20px}li{margin-bottom:5px}
blockquote{border-left:3px solid #6366f1;margin:14px 0;padding:8px 14px;background:#f0f0ff;border-radius:0 6px 6px 0}
code{font-family:'SF Mono',monospace;font-size:11px;background:#f0f0f5;padding:1px 4px;border-radius:3px}
pre{background:#f0f0f5;border:1px solid #d0d0e0;border-radius:6px;padding:10px;overflow-x:auto}
pre code{background:none;padding:0}
.header{text-align:center;margin-bottom:24px;padding-bottom:16px;border-bottom:2px solid #e0e0f0}
.meta{font-size:11px;color:#7a7a9e;margin-top:6px}
.cover{text-align:center;padding:60px 0 40px}
.cover h1{font-size:28px;border:none;margin-bottom:10px}
.cover .subtitle{font-size:14px;color:#666;margin-bottom:30px}
.cover .info{font-size:12px;color:#888;line-height:2}
"""

_PREVIEW_STYLE = """
body{font-family:'Segoe UI',Arial,sans-serif;max-width:800px;margin:40px auto;padding:0 30px;color:#1a1a2e;line-height:1.7}
h1{font-size:26px;border-bottom:3px solid #6366f1;padding-bottom:10px}
h2{font-size:20px;color:#2d2d44;margin-top:32px;border-left:4px solid #6366f1;padding-left:12px}
h3{font-size:16px;color:#3a3a52;margin-top:24px}
table{width:100%;border-collapse:collapse;margin:16px 0;font-size:13px}
th{background:#f0f0f5;padding:10px 12px;text-align:left;border:1px solid #d0d0e0;font-weight:600}
td{padding:8px 12px;border:1px solid #d0d0e0}
tr:nth-child(even){background:#f8f8fc}
ul,ol{padding-left:22px}li{margin-bottom:6px}
blockquote{border-left:3px solid #6366f1;margin:16px 0;padding:10px 16px;background:#f0f0ff}
.header{text-align:center;margin-bottom:30px;padding-bottom:20px;border-bottom:2px solid #e0e0f0}
.meta{font-size:12px;color:#7a7a9e;margin-top:8px}
"""


def build_html(content: str, project: dict, doc: dict, blueprint: dict,
               provider: str, watermark: dict | None = None, *, pdf: bool) -> str:
    # markdown 3.x ignores table/fenced_code kwargs passed to Markdown(...)
    md = markdown.Markdown(extensions=['tables', 'fenced_code'])
    body_html = md.convert(content)
    title = blueprint.get('projectTitle', project['name'])

    if pdf:
        wm_text = watermark.get("text", "") if watermark else ""
        wm_style = ""
        if wm_text:
            opacity = watermark.get("opacity", 0.15)
            font_size = watermark.get("fontSize", 48)
            wm_style = f"""
        body::after {{
            content: "{wm_text}";
            position: fixed;
            top: 50%; left: 50%;
            transform: translate(-50%, -50%) rotate(-35deg);
            font-size: {font_size}px;
            color: rgba(128,128,128,{opacity});
            pointer-events: none;
            z-index: 999;
            white-space: nowrap;
            font-weight: bold;
        }}"""
        style = _PDF_STYLE.replace("__TITLE__", title) + wm_style
        intro = f"""<div class="cover">
    <h1>{title}</h1>
    <div class="subtitle">{doc['name']}</div>
    <div class="info">
        {project['sector']} | {project['geo']} | {project['client']}<br>
        Generated via {provider} | {time.strftime('%B %d, %Y')}
    </div>
</div>
<div style="page-break-before:always"></div>"""
    else:
        style = _PREVIEW_STYLE
        intro = f"""<div class="header"><h1>{title}</h1>
<p class="meta">{doc['name']} | {project['sector']} | {project['geo']} | via {provider}</p></div>"""

    return f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><style>
{style}
</style></head><body>
{intro}
{body_html}
</body></html>"""


def render_pdf(content: str, project: dict, doc: dict, blueprint: dict,
               provider: str, watermark: dict | None = None) -> bytes:
    html = build_html(content, project, doc, blueprint, provider, watermark, pdf=True)
    return WeasyHTML(string=html).write_pdf()


def render_docx(content: str, project: dict, doc: dict, blueprint: dict,
                provider: str, watermark: dict | None = None) -> bytes:
    d = Document()

    section = d.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    wm_text = watermark.get("text", "") if watermark else ""
    if wm_text:
        for para in d.sections[0].header.paragraphs:
            run = para.add_run(wm_text)
            run.font.size = Pt(watermark.get("fontSize", 36))
            run.font.color.rgb = RGBColor(180, 180, 180)
            run.font.bold = True
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        section.header_distance = Cm(1)

    d.add_paragraph()
    d.add_paragraph()
    title = d.add_heading(blueprint.get('projectTitle', project['name']), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = d.add_paragraph(doc['name'])
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor(100, 100, 100)

    d.add_paragraph()

    info = d.add_paragraph(f"{project['sector']} | {project['geo']} | {project['client']}")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.runs[0].font.size = Pt(11)
    info.runs[0].font.color.rgb = RGBColor(120, 120, 120)

    gen_info = d.add_paragraph(f"Generated via {provider} | {time.strftime('%B %d, %Y')}")
    gen_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gen_info.runs[0].font.size = Pt(10)
    gen_info.runs[0].font.color.rgb = RGBColor(150, 150, 150)

    d.add_page_break()

    in_table = False
    table_rows: list = []

    for line in content.split('\n'):
        stripped = line.strip()

        if stripped.startswith('### '):
            d.add_heading(stripped[4:], level=3)
        elif stripped.startswith('## '):
            d.add_heading(stripped[3:], level=2)
        elif stripped.startswith('# '):
            d.add_heading(stripped[2:], level=1)
        elif stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            if all(set(c) <= set('- :') for c in cells):
                continue
            table_rows.append(cells)
            in_table = True
        else:
            if in_table and table_rows:
                _add_table_to_docx(d, table_rows)
                table_rows, in_table = [], False

            if stripped.startswith('- ') or stripped.startswith('* '):
                d.add_paragraph(stripped[2:], style='List Bullet')
            elif stripped.startswith('1. ') or stripped.startswith('2. '):
                d.add_paragraph(stripped[3:], style='List Number')
            elif stripped.startswith('> '):
                p = d.add_paragraph(stripped[2:])
                p.paragraph_format.left_indent = Cm(1)
                if p.runs:
                    p.runs[0].font.italic = True
                    p.runs[0].font.color.rgb = RGBColor(100, 100, 100)
            elif stripped in ('---', '***', '___'):
                p = d.add_paragraph('─' * 60)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if p.runs:
                    p.runs[0].font.color.rgb = RGBColor(200, 200, 200)
            elif not stripped:
                pass
            else:
                d.add_paragraph(stripped)

    if table_rows:
        _add_table_to_docx(d, table_rows)

    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()


def _add_table_to_docx(d, rows: list) -> None:
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    table = d.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j < num_cols:
                cell = table.cell(i, j)
                cell.text = cell_text
                if i == 0:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.bold = True
                            run.font.size = Pt(11)
